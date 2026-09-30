#!/usr/bin/env python3
"""CASH-AI: reproduzierbare, leakage-sichere Datenaufbereitung.

Das Programm liest die sechs Rohdatentabellen des Projekts, bildet getrennte
Modelltabellen für Kundeneinzahlungen und Beschaffungsauszahlungen, erzeugt
tägliche und wöchentliche Cashflow-Zeitreihen und teilt alle Daten rein
chronologisch in Training und Validierung. Es trainiert bewusst kein Modell.

Getestet mit Python 3.11/3.12 sowie den Dummy-Daten mit Stichtag 30.06.2026.
Die Berechnungslogik ist für Dummy- und Echtdaten identisch. Für Echtdaten sind
nur CONFIG bzw. die Kommandozeilenargumente anzupassen.
"""

from __future__ import annotations

import argparse
import bisect
import csv
import hashlib
import json
import logging
import math
import os
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

os.environ.setdefault("MPLCONFIGDIR", str(Path(__file__).resolve().parent / ".matplotlib_cache"))

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


# ---------------------------------------------------------------------------
# 1. ZENTRALE KONFIGURATION – für Echtdaten nur hier/über CLI anpassen
# ---------------------------------------------------------------------------

PROGRAMMORDNER = Path(__file__).resolve().parent

CONFIG: dict[str, Any] = {
    "input_folder": PROGRAMMORDNER / "Eingabedaten",
    "output_folder": PROGRAMMORDNER / "Ergebnisse_CASH_AI",
    "data_mode": "dummy",                 # "dummy" oder "real"
    "datenstichtag": "2026-06-30",
    "historie_von": None,                # automatisch aus den Rohdaten ermitteln
    "validierungstage": 42,
    "forecast_horizons": [7, 14, 30],
    "rolling_windows": [7, 14, 30, 90],
    "date_order": "monthfirst",           # aktuelle Dummy-Exporte; real oft "dayfirst"
    "min_history": 3,
    "source_files": {
        "kundenstammdaten": "kundenstammdaten.csv",
        "angebote": "angebote.csv",
        "ausgangsrechnungen": "ausgangsrechnungen.csv",
        "lieferantenstammdaten": "lieferantenstammdaten.csv",
        "bestellungen": "bestellungen.csv",
        "lagerbestand": "lagerbestand.csv",
    },
}


PRIMARY_KEYS = {
    "kundenstammdaten": "kunden_id",
    "angebote": "angebots_id",
    "ausgangsrechnungen": "rechnungs_id",
    "lieferantenstammdaten": "lieferanten_id",
    "bestellungen": "bestellungs_id",
}

REQUIRED_COLUMNS = {
    "kundenstammdaten": [
        "kunden_id", "kundenname", "segment", "region", "kundenklasse",
        "kunde_seit", "zahlungsziel_tage", "skonto_konditionen", "bonitaetsklasse",
    ],
    "angebote": [
        "angebots_id", "kunden_id", "angebotsdatum", "maschinenkategorie",
        "angebotswert", "entscheidungsdatum", "status", "bestellungs_id", "rechnungs_id",
    ],
    "ausgangsrechnungen": [
        "rechnungs_id", "kunden_id", "rechnungsdatum", "geschaeftsbereich",
        "materialkategorie", "angebots_id", "rechnungsbetrag", "zahlungsziel_tage",
        "faelligkeitsdatum", "skontofrist_tage", "skontosatz", "zahlungsdatum_ist",
        "bezahlter_betrag", "zahlungsstatus",
    ],
    "lieferantenstammdaten": [
        "lieferanten_id", "lieferantenname", "kategorie", "zahlungsziel_tage",
        "skontofrist_tage", "skontosatz",
    ],
    "bestellungen": [
        "bestellungs_id", "lieferanten_id", "geschaeftsbereich", "materialkategorie",
        "angebots_id", "bestelldatum", "bestellwert", "erwarteter_wareneingang",
        "tatsaechlicher_wareneingang", "zahlungsziel_tage", "faelligkeitsdatum",
        "skontofrist_tage", "skontosatz", "zahlungsdatum_ist", "bezahlter_betrag",
    ],
    "lagerbestand": [
        "monat", "kategorie", "bestand_anfang", "zugang", "abgang",
        "bestand_ende", "meldebestand",
    ],
}

DATE_FIELDS = {
    "kundenstammdaten": ["kunde_seit"],
    "angebote": ["angebotsdatum", "entscheidungsdatum"],
    "ausgangsrechnungen": ["rechnungsdatum", "faelligkeitsdatum", "zahlungsdatum_ist"],
    "bestellungen": [
        "bestelldatum", "erwarteter_wareneingang", "tatsaechlicher_wareneingang",
        "faelligkeitsdatum", "zahlungsdatum_ist",
    ],
    "lagerbestand": ["monat"],
}

NUMERIC_FIELDS = {
    "kundenstammdaten": ["zahlungsziel_tage"],
    "angebote": ["angebotswert"],
    "ausgangsrechnungen": [
        "rechnungsbetrag", "zahlungsziel_tage", "skontofrist_tage", "skontosatz",
        "bezahlter_betrag",
    ],
    "lieferantenstammdaten": ["zahlungsziel_tage", "skontofrist_tage", "skontosatz"],
    "bestellungen": [
        "bestellwert", "zahlungsziel_tage", "skontofrist_tage", "skontosatz",
        "bezahlter_betrag",
    ],
    "lagerbestand": ["bestand_anfang", "zugang", "abgang", "bestand_ende", "meldebestand"],
}

# Über diese Aliasliste lassen sich Echtdatenfelder zentral zuordnen.
COLUMN_ALIASES: dict[str, list[str]] = {
    "kunden_id": ["kundennummer", "kunde_id", "customer_id", "debitorennummer"],
    "kundenname": ["kunde", "customer_name", "debitorenname"],
    "kunde_seit": ["kund_seit", "customer_since"],
    "kundenklasse": ["kundengruppe", "customer_class"],
    "bonitaetsklasse": ["bonitaet", "credit_class"],
    "lieferanten_id": ["lieferantennummer", "vendor_id", "kreditorennummer"],
    "lieferantenname": ["lieferant", "vendor_name", "kreditorenname"],
    "angebots_id": ["angebotsnummer", "angebot_id", "quote_id"],
    "angebotsdatum": ["angebot_datum", "quote_date"],
    "angebotswert": ["angebotssumme", "quote_amount"],
    "entscheidungsdatum": ["entscheidung_datum", "decision_date"],
    "bestellungs_id": ["bestellnummer", "bestellung_id", "purchase_order_id"],
    "bestelldatum": ["bestellung_datum", "order_date"],
    "bestellwert": ["bestellsumme", "order_amount"],
    "rechnungs_id": ["rechnungsnummer", "rechnung_id", "invoice_id"],
    "rechnungsdatum": ["rechnung_datum", "invoice_date"],
    "rechnungsbetrag": ["rechnungssumme", "invoice_amount", "bruttobetrag"],
    "faelligkeitsdatum": ["faellig_am", "due_date"],
    "zahlungsdatum_ist": ["zahlungsdatum", "payment_date", "auszahlungsdatum"],
    "bezahlter_betrag": ["zahlbetrag", "payment_amount", "auszahlungsbetrag"],
    "zahlungsziel_tage": ["zahlungsziel", "payment_terms_days"],
    "skontofrist_tage": ["skontofrist", "discount_days"],
    "skontosatz": ["skonto_prozent", "discount_rate"],
    "erwarteter_wareneingang": ["geplanter_wareneingang", "expected_receipt_date"],
    "tatsaechlicher_wareneingang": ["wareneingangsdatum", "actual_receipt_date"],
    "geschaeftsbereich": ["bereich", "business_area"],
    "materialkategorie": ["materialgruppe", "material_category"],
    "maschinenkategorie": ["maschinenart", "machine_category"],
}

WEEKDAY_DE = {
    0: "Montag", 1: "Dienstag", 2: "Mittwoch", 3: "Donnerstag",
    4: "Freitag", 5: "Samstag", 6: "Sonntag",
}


# ---------------------------------------------------------------------------
# 2. HILFSFUNKTIONEN: Konfiguration, Logging, sichere Konvertierung
# ---------------------------------------------------------------------------

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="CASH-AI Datenaufbereitung (noch ohne Modelltraining)"
    )
    parser.add_argument("--input-folder", help="Ordner mit den sechs Rohdateien")
    parser.add_argument("--output-folder", help="Zielordner für alle Ergebnisse")
    parser.add_argument("--data-mode", choices=["dummy", "real"])
    parser.add_argument("--datenstichtag", help="Stichtag im Format YYYY-MM-DD")
    parser.add_argument("--validierungstage", type=int)
    parser.add_argument("--date-order", choices=["monthfirst", "dayfirst", "auto"])
    return parser.parse_args()


def effective_config(args: argparse.Namespace) -> dict[str, Any]:
    cfg = dict(CONFIG)
    cfg["source_files"] = dict(CONFIG["source_files"])
    for arg_name, key in [
        ("input_folder", "input_folder"), ("output_folder", "output_folder"),
        ("data_mode", "data_mode"), ("datenstichtag", "datenstichtag"),
        ("validierungstage", "validierungstage"), ("date_order", "date_order"),
    ]:
        value = getattr(args, arg_name)
        if value is not None:
            cfg[key] = value
    cfg["input_folder"] = Path(cfg["input_folder"]).expanduser().resolve()
    cfg["output_folder"] = Path(cfg["output_folder"]).expanduser().resolve()
    cfg["datenstichtag"] = pd.Timestamp(cfg["datenstichtag"]).normalize()
    if cfg.get("historie_von"):
        cfg["historie_von"] = pd.Timestamp(cfg["historie_von"]).normalize()
    if int(cfg["validierungstage"]) < 7:
        raise ValueError("validierungstage muss mindestens 7 betragen.")
    cfg["validierungstage"] = int(cfg["validierungstage"])
    cfg["validierung_start"] = cfg["datenstichtag"] - pd.Timedelta(
        days=cfg["validierungstage"] - 1
    )
    cfg["training_ende"] = cfg["validierung_start"] - pd.Timedelta(days=1)
    return cfg


def setup_logging(output_folder: Path) -> Path:
    output_folder.mkdir(parents=True, exist_ok=True)
    log_path = output_folder / "14_CASH_AI_Verarbeitungsprotokoll.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s",
        handlers=[
            logging.FileHandler(log_path, mode="w", encoding="utf-8"),
            logging.StreamHandler(sys.stdout),
        ],
        force=True,
    )
    return log_path


def normalize_name(value: str) -> str:
    value = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "_", value.lower()).strip("_")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def json_safe(value: Any) -> Any:
    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, pd.Timestamp):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating, float)):
        return None if not math.isfinite(float(value)) else float(value)
    if isinstance(value, (np.bool_, bool)):
        return bool(value)
    if isinstance(value, dict):
        return {str(key): json_safe(val) for key, val in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_safe(val) for val in value]
    return value


def format_keys(values: Iterable[Any], limit: int = 25) -> str:
    keys = [str(x) for x in values if pd.notna(x)]
    if len(keys) > limit:
        return ", ".join(keys[:limit]) + f" … (+{len(keys) - limit})"
    return ", ".join(keys)


def add_check(
    rows: list[dict[str, Any]], category: str, dataset: str, name: str,
    description: str, violations: int, keys: Iterable[Any] = (),
    severity: str = "kritisch",
) -> None:
    if violations == 0:
        status = "PASS"
    elif severity == "warnung":
        status = "WARN"
    else:
        status = "FAIL"
    rows.append({
        "Kategorie": category,
        "Datensatz": dataset,
        "Pruefung": name,
        "Beschreibung": description,
        "Verletzungen": int(violations),
        "Status": status,
        "Schweregrad": severity,
        "Betroffene_Schluessel": format_keys(keys),
    })


def parse_number(value: Any) -> float:
    if value is None or pd.isna(value):
        return np.nan
    text = str(value).strip().replace("\u00a0", "").replace(" ", "")
    text = re.sub(r"[€$%]", "", text)
    if text == "" or text.lower() in {"nan", "none", "null", "na", "n/a"}:
        return np.nan
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return np.nan


def parse_date_series(series: pd.Series, date_order: str) -> pd.Series:
    """ISO zuerst; nicht-ISO gemäß zentraler Datumsreihenfolge verarbeiten."""
    raw = series.astype("string").str.strip()
    result = pd.Series(pd.NaT, index=series.index, dtype="datetime64[ns]")
    empty = raw.isna() | raw.str.lower().isin(["", "nan", "none", "null", "nat"])
    iso_mask = raw.str.match(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}$", na=False)
    result.loc[iso_mask] = pd.to_datetime(
        raw.loc[iso_mask].str.replace("/", "-", regex=False), format="%Y-%m-%d", errors="coerce"
    )
    remaining = ~empty & result.isna()
    if date_order == "dayfirst":
        formats = ["%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%y", "%d/%m/%y"]
    else:
        formats = ["%m.%d.%Y", "%m/%d/%Y", "%m-%d-%Y", "%m.%d.%y", "%m/%d/%y"]
    for fmt in formats:
        if not remaining.any():
            break
        parsed = pd.to_datetime(raw.loc[remaining], format=fmt, errors="coerce")
        valid_index = parsed.index[parsed.notna()]
        result.loc[valid_index] = parsed.loc[valid_index]
        remaining = ~empty & result.isna()
    if remaining.any() and date_order == "auto":
        result.loc[remaining] = pd.to_datetime(raw.loc[remaining], errors="coerce")
    return result.dt.normalize()


def detect_separator(path: Path) -> str:
    sample = path.read_text(encoding="utf-8-sig", errors="replace")[:16384]
    try:
        return csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
    except csv.Error:
        return ";" if sample.count(";") > sample.count(",") else ","


def resolve_source_file(input_folder: Path, configured_name: str) -> Path:
    exact = input_folder / configured_name
    if exact.exists():
        return exact
    stem = Path(configured_name).stem
    suffix = Path(configured_name).suffix or ".csv"
    pattern = re.compile(rf"^{re.escape(stem)}(?:\((\d+)\))?{re.escape(suffix)}$", re.IGNORECASE)
    candidates: list[tuple[int, Path]] = []
    for path in input_folder.iterdir():
        match = pattern.match(path.name)
        if match and path.is_file():
            candidates.append((int(match.group(1) or 0), path))
    if not candidates:
        raise FileNotFoundError(
            f"Quelldatei '{configured_name}' fehlt in '{input_folder}'. "
            "Bitte Dateiname in CONFIG['source_files'] anpassen."
        )
    chosen = sorted(candidates, key=lambda item: item[0])[-1][1]
    logging.warning("%s nicht exakt gefunden; verwende %s", configured_name, chosen.name)
    return chosen


def map_columns(df: pd.DataFrame, table: str) -> tuple[pd.DataFrame, list[dict[str, Any]]]:
    normalized_actual = {normalize_name(col): col for col in df.columns}
    rename: dict[str, str] = {}
    provenance: list[dict[str, Any]] = []
    missing: list[str] = []
    for canonical in REQUIRED_COLUMNS[table]:
        aliases = [canonical] + COLUMN_ALIASES.get(canonical, [])
        matches = [normalized_actual[normalize_name(alias)] for alias in aliases if normalize_name(alias) in normalized_actual]
        matches = list(dict.fromkeys(matches))
        if len(matches) == 0:
            missing.append(canonical)
            continue
        if len(matches) > 1:
            raise ValueError(
                f"{table}: Feld '{canonical}' ist nicht eindeutig zuordenbar: {matches}. "
                "Bitte COLUMN_ALIASES bereinigen."
            )
        rename[matches[0]] = canonical
        provenance.append({
            "Datensatz": table, "Ergebnisfeld": canonical,
            "Quellspalte": matches[0], "Art": "Direkte Zuordnung",
        })
    if missing:
        raise ValueError(
            f"{table}: Pflichtspalten fehlen: {missing}. "
            "Bitte COLUMN_ALIASES oder die Quelldatei anpassen."
        )
    out = df.rename(columns=rename)
    return out[REQUIRED_COLUMNS[table]].copy(), provenance


# ---------------------------------------------------------------------------
# 3. QUELLEN EINLESEN UND SCHEMA/DATENQUALITÄT PRÜFEN
# ---------------------------------------------------------------------------

def inspect_context_files(input_folder: Path) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for path in sorted(input_folder.iterdir()):
        if not path.is_file():
            continue
        row = {
            "Datei": path.name,
            "Typ": path.suffix.lower(),
            "Bytes": path.stat().st_size,
            "Verwendung": "Kontext/Dokumentation" if path.suffix.lower() in {".xlsx", ".xls", ".md", ".py"} else "Rohdatenkandidat",
            "Tabellenblaetter": "",
        }
        if path.suffix.lower() == ".xlsx":
            try:
                workbook = load_workbook(path, read_only=True, data_only=False)
                row["Tabellenblaetter"] = ", ".join(workbook.sheetnames)
                workbook.close()
            except Exception as exc:  # Kontextdateien dürfen die Pipeline nicht blockieren.
                row["Tabellenblaetter"] = f"Nicht lesbar: {exc}"
        rows.append(row)
    return pd.DataFrame(rows)


def load_sources(
    cfg: dict[str, Any], checks: list[dict[str, Any]]
) -> tuple[dict[str, pd.DataFrame], pd.DataFrame, pd.DataFrame]:
    input_folder = cfg["input_folder"]
    if not input_folder.exists():
        raise FileNotFoundError(
            f"Eingabeordner fehlt: {input_folder}. Lege dort die sechs Roh-CSV-Dateien ab."
        )
    frames: dict[str, pd.DataFrame] = {}
    inventory_rows: list[dict[str, Any]] = []
    provenance_rows: list[dict[str, Any]] = []
    for table, configured_name in cfg["source_files"].items():
        path = resolve_source_file(input_folder, configured_name)
        separator = detect_separator(path)
        raw = pd.read_csv(
            path, sep=separator, dtype=str, encoding="utf-8-sig", keep_default_na=True,
            na_values=["", "NA", "N/A", "NULL", "null", "None"], low_memory=False,
        )
        mapped, provenance = map_columns(raw, table)
        provenance_rows.extend(provenance)
        for column in DATE_FIELDS.get(table, []):
            original = mapped[column].copy()
            raw_nonempty = int(original.notna().sum())
            converted = parse_date_series(original, cfg["date_order"])
            invalid = raw_nonempty - int(converted.notna().sum())
            mapped[column] = converted
            invalid_mask = original.notna() & mapped[column].isna()
            bad_keys = mapped.loc[invalid_mask, PRIMARY_KEYS.get(table, mapped.columns[0])]
            add_check(
                checks, "Datentyp", table, f"Datumsformat {column}",
                "Nichtleere Datumswerte müssen konvertierbar sein.", invalid, bad_keys,
                "kritisch" if column in {"rechnungsdatum", "bestelldatum"} else "warnung",
            )
        for column in NUMERIC_FIELDS.get(table, []):
            original = mapped[column].copy()
            mapped[column] = original.map(parse_number)
            invalid_mask = original.notna() & mapped[column].isna()
            add_check(
                checks, "Datentyp", table, f"Zahlenformat {column}",
                "Nichtleere Zahlenwerte müssen konvertierbar sein.", int(invalid_mask.sum()),
                mapped.loc[invalid_mask, PRIMARY_KEYS.get(table, mapped.columns[0])],
                "kritisch" if column in {"rechnungsbetrag", "bestellwert"} else "warnung",
            )
        pk = PRIMARY_KEYS.get(table)
        if pk:
            missing_pk = mapped[pk].isna() | mapped[pk].astype(str).str.strip().eq("")
            duplicate = mapped[pk].duplicated(keep=False) & ~missing_pk
            add_check(checks, "Schluessel", table, f"Pflichtfeld {pk}", "Primärschlüssel darf nicht fehlen.", int(missing_pk.sum()), mapped.loc[missing_pk, pk])
            add_check(checks, "Schluessel", table, f"Eindeutigkeit {pk}", "Primärschlüssel muss eindeutig sein.", int(duplicate.sum()), mapped.loc[duplicate, pk])
        frames[table] = mapped
        inventory_rows.append({
            "Datensatz": table, "Quelldatei": path.name, "Pfad": str(path),
            "Trennzeichen": repr(separator), "Zeilen": len(mapped), "Spalten": len(mapped.columns),
            "SHA256": sha256(path),
        })
        logging.info("Quelle geladen: %s (%s Zeilen, %s Spalten)", path.name, len(mapped), len(mapped.columns))
    return frames, pd.DataFrame(inventory_rows), pd.DataFrame(provenance_rows)


def validate_sources(d: dict[str, pd.DataFrame], checks: list[dict[str, Any]]) -> None:
    relationships = [
        ("angebote", "kunden_id", "kundenstammdaten", "kunden_id"),
        ("ausgangsrechnungen", "kunden_id", "kundenstammdaten", "kunden_id"),
        ("bestellungen", "lieferanten_id", "lieferantenstammdaten", "lieferanten_id"),
        ("ausgangsrechnungen", "angebots_id", "angebote", "angebots_id"),
        ("bestellungen", "angebots_id", "angebote", "angebots_id"),
    ]
    for child, fk, parent, pk in relationships:
        values = d[child][fk].dropna()
        invalid = ~values.isin(set(d[parent][pk].dropna()))
        add_check(
            checks, "Beziehungen", child, f"Fremdschluessel {fk} -> {parent}.{pk}",
            "Jeder befüllte Fremdschlüssel muss in der Stammtabelle vorkommen.",
            int(invalid.sum()), values.loc[invalid], "kritisch",
        )

    invoices = d["ausgangsrechnungen"]
    orders = d["bestellungen"]
    invalid_payment = invoices["zahlungsdatum_ist"].notna() & (
        invoices["zahlungsdatum_ist"] < invoices["rechnungsdatum"]
    )
    invalid_due = invoices["faelligkeitsdatum"].notna() & (
        invoices["faelligkeitsdatum"] < invoices["rechnungsdatum"]
    )
    invalid_receipt = orders["tatsaechlicher_wareneingang"].notna() & (
        orders["tatsaechlicher_wareneingang"] < orders["bestelldatum"]
    )
    invalid_order_payment = orders["zahlungsdatum_ist"].notna() & (
        orders["zahlungsdatum_ist"] < orders["bestelldatum"]
    )
    add_check(checks, "Plausibilitaet", "ausgangsrechnungen", "Zahlung vor Rechnungsdatum", "Zeilen bleiben erhalten, werden aber aus Zahlungshistorien ausgeschlossen.", int(invalid_payment.sum()), invoices.loc[invalid_payment, "rechnungs_id"], "warnung")
    add_check(checks, "Plausibilitaet", "ausgangsrechnungen", "Faelligkeit vor Rechnungsdatum", "Fälligkeitsdatum darf nicht vor Rechnungsdatum liegen.", int(invalid_due.sum()), invoices.loc[invalid_due, "rechnungs_id"], "warnung")
    add_check(checks, "Plausibilitaet", "bestellungen", "Wareneingang vor Bestelldatum", "Zeilen bleiben erhalten, werden aber aus Lieferhistorien ausgeschlossen.", int(invalid_receipt.sum()), orders.loc[invalid_receipt, "bestellungs_id"], "warnung")
    add_check(checks, "Plausibilitaet", "bestellungen", "Zahlung vor Bestelldatum", "Zeilen bleiben erhalten, werden aber aus Zahlungshistorien ausgeschlossen.", int(invalid_order_payment.sum()), orders.loc[invalid_order_payment, "bestellungs_id"], "warnung")

    for table, column in [
        ("ausgangsrechnungen", "rechnungsbetrag"), ("ausgangsrechnungen", "bezahlter_betrag"),
        ("bestellungen", "bestellwert"), ("bestellungen", "bezahlter_betrag"),
    ]:
        negative = d[table][column].notna() & d[table][column].lt(0)
        add_check(checks, "Plausibilitaet", table, f"Negative Betraege {column}", "Beträge dürfen nicht negativ sein.", int(negative.sum()), d[table].loc[negative, PRIMARY_KEYS[table]], "warnung")


# ---------------------------------------------------------------------------
# 4. FEATURE ENGINEERING – strikt nur Informationen vor dem Stichtag
# ---------------------------------------------------------------------------

def easter_sunday(year: int) -> pd.Timestamp:
    a = year % 19
    b, c = divmod(year, 100)
    d, e = divmod(b, 4)
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = divmod(c, 4)
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    month = (h + l - 7 * m + 114) // 31
    day = ((h + l - 7 * m + 114) % 31) + 1
    return pd.Timestamp(year=year, month=month, day=day)


def austrian_holidays(years: Iterable[int]) -> dict[pd.Timestamp, str]:
    holidays: dict[pd.Timestamp, str] = {}
    fixed = {
        (1, 1): "Neujahr", (1, 6): "Heilige Drei Koenige", (5, 1): "Staatsfeiertag",
        (8, 15): "Mariae Himmelfahrt", (10, 26): "Nationalfeiertag",
        (11, 1): "Allerheiligen", (12, 8): "Mariae Empfaengnis",
        (12, 25): "Christtag", (12, 26): "Stefanitag",
    }
    for year in years:
        for (month, day), name in fixed.items():
            holidays[pd.Timestamp(year=year, month=month, day=day)] = name
        easter = easter_sunday(year)
        holidays[easter + pd.Timedelta(days=1)] = "Ostermontag"
        holidays[easter + pd.Timedelta(days=39)] = "Christi Himmelfahrt"
        holidays[easter + pd.Timedelta(days=50)] = "Pfingstmontag"
        holidays[easter + pd.Timedelta(days=60)] = "Fronleichnam"
    return holidays


def add_calendar_features(df: pd.DataFrame, column: str, prefix: str, holidays: set[pd.Timestamp]) -> None:
    dates = df[column]
    iso = dates.dt.isocalendar()
    df[f"{prefix}_wochentag_nr"] = (dates.dt.weekday + 1).astype("Int64")
    df[f"{prefix}_wochentag"] = dates.dt.weekday.map(WEEKDAY_DE)
    df[f"{prefix}_monat"] = dates.dt.month.astype("Int64")
    df[f"{prefix}_quartal"] = dates.dt.quarter.astype("Int64")
    df[f"{prefix}_kalenderwoche"] = iso.week.astype("Int64")
    df[f"{prefix}_jahr"] = dates.dt.year.astype("Int64")
    df[f"{prefix}_tag_im_monat"] = dates.dt.day.astype("Int64")
    df[f"{prefix}_wochenende"] = dates.dt.weekday.ge(5).astype("Int64")
    df[f"{prefix}_feiertag_at"] = dates.map(lambda value: int(pd.Timestamp(value) in holidays) if pd.notna(value) else pd.NA).astype("Int64")
    df[f"{prefix}_feiertagsnaehe_3t"] = dates.map(
        lambda value: int(any((pd.Timestamp(value) + pd.Timedelta(days=delta)) in holidays for delta in range(-3, 4))) if pd.notna(value) else pd.NA
    ).astype("Int64")


def historical_transaction_windows(
    frame: pd.DataFrame, entity: str, date_col: str, amount_col: str,
    windows: Iterable[int], prefix: str,
) -> pd.DataFrame:
    """Rollierende Mengen/Beträge; gleiche Kalendertage werden ausgeschlossen."""
    base = frame[[entity, date_col, amount_col]].copy()
    base["_row"] = frame.index
    daily = base.groupby([entity, date_col], as_index=False).agg(
        _betrag=(amount_col, "sum"), _anzahl=(amount_col, "size")
    )
    pieces: list[pd.DataFrame] = []
    for entity_value, group in daily.groupby(entity, sort=False):
        group = group.sort_values(date_col).set_index(date_col)
        result = pd.DataFrame({entity: entity_value, date_col: group.index})
        for window in windows:
            result[f"{prefix}_volumen_letzte_{window}t"] = group["_betrag"].rolling(f"{window}D", closed="left").sum().to_numpy()
            result[f"{prefix}_anzahl_letzte_{window}t"] = group["_anzahl"].rolling(f"{window}D", closed="left").sum().to_numpy()
        pieces.append(result)
    rolled = pd.concat(pieces, ignore_index=True) if pieces else pd.DataFrame()
    return base.merge(rolled, on=[entity, date_col], how="left", validate="many_to_one").set_index("_row")[[
        col for col in rolled.columns if col not in {entity, date_col}
    ]]


def build_neumaschinen_daily(
    d: dict[str, pd.DataFrame], cfg: dict[str, Any]
) -> pd.DataFrame:
    history_start = cfg["historie_von"] or min(
        d["ausgangsrechnungen"]["rechnungsdatum"].min(), d["bestellungen"]["bestelldatum"].min()
    )
    index = pd.date_range(history_start, cfg["datenstichtag"], freq="D")
    invoices = d["ausgangsrechnungen"]
    orders = d["bestellungen"]
    offers = d["angebote"]
    out = pd.DataFrame(index=index)
    out["ziel_nm_verkaufsvolumen_tag"] = invoices.loc[
        invoices["geschaeftsbereich"].eq("Neumaschinen")
    ].groupby("rechnungsdatum")["rechnungsbetrag"].sum().reindex(index, fill_value=0.0)
    out["ziel_nm_bestellvolumen_tag"] = orders.loc[
        orders["geschaeftsbereich"].eq("Neumaschinen")
    ].groupby("bestelldatum")["bestellwert"].sum().reindex(index, fill_value=0.0)
    won = offers.loc[offers["status"].astype(str).str.lower().eq("gewonnen") & offers["entscheidungsdatum"].notna()]
    out["ziel_nm_gewonnene_angebote_anzahl_tag"] = won.groupby("entscheidungsdatum").size().reindex(index, fill_value=0).astype(int)
    out["ziel_nm_gewonnene_angebote_wert_tag"] = won.groupby("entscheidungsdatum")["angebotswert"].sum().reindex(index, fill_value=0.0)
    for window in [7, 14, 30]:
        sales = out["ziel_nm_verkaufsvolumen_tag"].shift(1).rolling(window, min_periods=1).sum()
        orders_roll = out["ziel_nm_bestellvolumen_tag"].shift(1).rolling(window, min_periods=1).sum()
        previous_sales = out["ziel_nm_verkaufsvolumen_tag"].shift(window + 1).rolling(window, min_periods=1).sum()
        previous_orders = out["ziel_nm_bestellvolumen_tag"].shift(window + 1).rolling(window, min_periods=1).sum()
        out[f"nm_verkaufsvolumen_letzte_{window}t"] = sales
        out[f"nm_bestellvolumen_letzte_{window}t"] = orders_roll
        out[f"nm_verhaeltnis_bestell_zu_verkauf_{window}t"] = orders_roll.div(sales.where(sales.ne(0)))
        out[f"nm_verhaeltnis_nicht_definiert_{window}t"] = sales.eq(0).astype(int)
        out[f"nm_verkaufsveraenderung_abs_{window}t"] = sales - previous_sales
        out[f"nm_bestellveraenderung_abs_{window}t"] = orders_roll - previous_orders
        out[f"nm_verkaufsveraenderung_pct_{window}t"] = (sales - previous_sales).div(previous_sales.where(previous_sales.ne(0)))
        out[f"nm_bestellveraenderung_pct_{window}t"] = (orders_roll - previous_orders).div(previous_orders.where(previous_orders.ne(0)))
        out[f"nm_gewonnene_angebote_anzahl_letzte_{window}t"] = out["ziel_nm_gewonnene_angebote_anzahl_tag"].shift(1).rolling(window, min_periods=1).sum()
        out[f"nm_gewonnene_angebote_wert_letzte_{window}t"] = out["ziel_nm_gewonnene_angebote_wert_tag"].shift(1).rolling(window, min_periods=1).sum()
    out.index.name = "berechnungsstichtag"
    return out.reset_index()


def customer_history_features(group: pd.DataFrame, min_history: int) -> pd.DataFrame:
    valid = group.loc[
        group["__zahlungsdatum_ist"].notna()
        & group["rechnungsdatum"].notna()
        & group["__zahlungsdatum_ist"].ge(group["rechnungsdatum"])
    ].copy()
    valid["_delay"] = (valid["__zahlungsdatum_ist"] - valid["faelligkeitsdatum"]).dt.days
    valid["_skonto"] = (valid["__bezahlter_betrag"] < valid["rechnungsbetrag"] - 0.005).astype(int)
    events = valid.sort_values(["__zahlungsdatum_ist", "rechnungs_id"]).to_dict("records")
    pointer = 0
    delays_sorted: list[float] = []
    delays: list[float] = []
    invoice_amounts: list[float] = []
    paid_amounts: list[float] = []
    late = ontime = discount = 0
    last_payment: pd.Timestamp | None = None
    rows: list[dict[str, Any]] = []
    for index, row in group.sort_values(["rechnungsdatum", "rechnungs_id"]).iterrows():
        cutoff = row["rechnungsdatum"]
        while pointer < len(events) and events[pointer]["__zahlungsdatum_ist"] < cutoff:
            event = events[pointer]
            delay = float(event["_delay"])
            bisect.insort(delays_sorted, delay)
            delays.append(delay)
            invoice_amounts.append(float(event["rechnungsbetrag"]))
            paid_amounts.append(float(event["__bezahlter_betrag"]))
            late += int(delay > 0)
            ontime += int(delay <= 0)
            discount += int(event["_skonto"])
            last_payment = event["__zahlungsdatum_ist"]
            pointer += 1
        count = len(delays)
        rows.append({
            "_index": index,
            "hist_bezahlte_rechnungen_anzahl": count,
            "hist_zahlungsverzoegerung_mittel_tage": float(np.mean(delays)) if count else np.nan,
            "hist_zahlungsverzoegerung_median_tage": float(np.median(delays_sorted)) if count else np.nan,
            "hist_zahlungsverzoegerung_std_tage": float(np.std(delays, ddof=0)) if count else np.nan,
            "hist_anteil_verspaetet": late / count if count else np.nan,
            "hist_anteil_puenktlich_oder_frueh": ontime / count if count else np.nan,
            "hist_rechnungsbetrag_mittel": float(np.mean(invoice_amounts)) if count else np.nan,
            "hist_rechnungsbetrag_median": float(np.median(invoice_amounts)) if count else np.nan,
            "hist_bezahlter_betrag_mittel": float(np.mean(paid_amounts)) if count else np.nan,
            "hist_skontonutzung_anteil": discount / count if count else np.nan,
            "hist_tage_seit_letzter_zahlung": int((cutoff - last_payment).days) if last_payment is not None else np.nan,
            "hist_letztes_zahlungsdatum": last_payment,
            "kunde_ohne_ausreichende_zahlungshistorie": int(count < min_history),
        })
    return pd.DataFrame(rows).set_index("_index")


def customer_open_items(group: pd.DataFrame) -> pd.DataFrame:
    valid_date = group["__zahlungsdatum_ist"].isna() | group["__zahlungsdatum_ist"].ge(group["rechnungsdatum"])
    base = group.loc[valid_date]
    excluded = group.loc[~valid_date]
    rows: list[dict[str, Any]] = []
    for index, row in group.sort_values(["rechnungsdatum", "rechnungs_id"]).iterrows():
        cutoff = row["rechnungsdatum"]
        open_mask = base["rechnungsdatum"].lt(cutoff) & (
            base["__zahlungsdatum_ist"].isna() | base["__zahlungsdatum_ist"].ge(cutoff)
        )
        open_items = base.loc[open_mask]
        overdue = open_items.loc[open_items["faelligkeitsdatum"].lt(cutoff)]
        age = (cutoff - overdue["faelligkeitsdatum"]).dt.days
        rows.append({
            "_index": index,
            "op_offene_rechnungen_anzahl": len(open_items),
            "op_offener_betrag": float(open_items["rechnungsbetrag"].sum()),
            "op_ueberfaellige_rechnungen_anzahl": len(overdue),
            "op_ueberfaelliger_betrag": float(overdue["rechnungsbetrag"].sum()),
            "op_ueberfaellig_1_7t_betrag": float(overdue.loc[age.between(1, 7), "rechnungsbetrag"].sum()),
            "op_ueberfaellig_8_14t_betrag": float(overdue.loc[age.between(8, 14), "rechnungsbetrag"].sum()),
            "op_ueberfaellig_15_30t_betrag": float(overdue.loc[age.between(15, 30), "rechnungsbetrag"].sum()),
            "op_ueberfaellig_ueber_30t_betrag": float(overdue.loc[age.gt(30), "rechnungsbetrag"].sum()),
            "op_max_ueberziehungstage": int(age.max()) if len(age) else 0,
            "op_dq_ausgeschlossene_rechnungen_anzahl": int((excluded["rechnungsdatum"] < cutoff).sum()),
        })
    return pd.DataFrame(rows).set_index("_index")


def build_customer_features(
    d: dict[str, pd.DataFrame], nm_daily: pd.DataFrame, cfg: dict[str, Any], holidays: set[pd.Timestamp]
) -> pd.DataFrame:
    invoices = d["ausgangsrechnungen"].copy().rename(columns={
        "zahlungsdatum_ist": "__zahlungsdatum_ist",
        "bezahlter_betrag": "__bezahlter_betrag",
        "zahlungsstatus": "__zahlungsstatus",
    })
    customers = d["kundenstammdaten"].rename(columns={
        "zahlungsziel_tage": "kunde_zahlungsziel_tage",
        "skonto_konditionen": "kunde_skonto_konditionen",
    })
    out = invoices.merge(customers, on="kunden_id", how="left", validate="many_to_one")
    out["berechnungsstichtag"] = out["rechnungsdatum"]
    add_calendar_features(out, "rechnungsdatum", "rechnung", holidays)
    add_calendar_features(out, "faelligkeitsdatum", "faelligkeit", holidays)
    out["tage_rechnung_bis_faelligkeit"] = (out["faelligkeitsdatum"] - out["rechnungsdatum"]).dt.days
    out["skontodatum"] = out["rechnungsdatum"] + pd.to_timedelta(out["skontofrist_tage"], unit="D")
    out["kundenbeziehung_tage"] = (out["rechnungsdatum"] - out["kunde_seit"]).dt.days
    out["dq_kunde_nicht_verknuepft"] = out["kundenname"].isna().astype(int)
    history = pd.concat([
        customer_history_features(group, cfg["min_history"])
        for _, group in out.groupby("kunden_id", sort=False)
    ]).sort_index()
    open_items = pd.concat([
        customer_open_items(group) for _, group in out.groupby("kunden_id", sort=False)
    ]).sort_index()
    rolling = historical_transaction_windows(
        out, "kunden_id", "rechnungsdatum", "rechnungsbetrag",
        cfg["rolling_windows"], "hist_rechnung",
    ).sort_index()
    out = out.join(history).join(open_items).join(rolling)
    safe_nm_columns = [
        col for col in nm_daily.columns
        if col == "berechnungsstichtag" or not col.startswith("ziel_")
    ]
    out = out.merge(nm_daily[safe_nm_columns], on="berechnungsstichtag", how="left", validate="many_to_one")
    out["ziel_dq_zahlung_vor_rechnungsdatum"] = (
        out["__zahlungsdatum_ist"].notna() & out["__zahlungsdatum_ist"].lt(out["rechnungsdatum"])
    ).astype(int)
    return out.sort_values(["rechnungsdatum", "rechnungs_id"]).reset_index(drop=True)


def supplier_history_features(group: pd.DataFrame, min_history: int) -> pd.DataFrame:
    valid = (
        group["__zahlungsdatum_ist"].notna()
        & group["__wareneingang_ist"].notna()
        & group["__zahlungsdatum_ist"].ge(group["bestelldatum"])
        & group["__wareneingang_ist"].ge(group["bestelldatum"])
    )
    completed = group.loc[valid].copy()
    completed["_abschluss"] = completed[["__zahlungsdatum_ist", "__wareneingang_ist"]].max(axis=1)
    completed["_zahlungsabweichung"] = (completed["__zahlungsdatum_ist"] - completed["faelligkeitsdatum"]).dt.days
    completed["_lieferzeit"] = (completed["__wareneingang_ist"] - completed["bestelldatum"]).dt.days
    completed["_wareneingangsabweichung"] = (completed["__wareneingang_ist"] - completed["erwarteter_wareneingang"]).dt.days
    completed["_skonto"] = (
        (completed["__zahlungsdatum_ist"] <= completed["bestelldatum"] + pd.to_timedelta(completed["skontofrist_tage"], unit="D"))
        & (completed["__bezahlter_betrag"] < completed["bestellwert"] - 0.005)
    ).astype(int)
    events = completed.sort_values(["_abschluss", "bestellungs_id"]).to_dict("records")
    pointer = 0
    amounts: list[float] = []
    payment_delta: list[float] = []
    delivery_time: list[float] = []
    receipt_delta: list[float] = []
    ontime = discount = 0
    last_payment: pd.Timestamp | None = None
    last_completion: pd.Timestamp | None = None
    rows: list[dict[str, Any]] = []
    for index, row in group.sort_values(["bestelldatum", "bestellungs_id"]).iterrows():
        cutoff = row["bestelldatum"]
        while pointer < len(events) and events[pointer]["_abschluss"] < cutoff:
            event = events[pointer]
            amounts.append(float(event["bestellwert"]))
            payment_delta.append(float(event["_zahlungsabweichung"]))
            delivery_time.append(float(event["_lieferzeit"]))
            receipt_delta.append(float(event["_wareneingangsabweichung"]))
            ontime += int(event["_zahlungsabweichung"] <= 0)
            discount += int(event["_skonto"])
            last_payment = event["__zahlungsdatum_ist"]
            last_completion = event["_abschluss"]
            pointer += 1
        count = len(amounts)
        rows.append({
            "_index": index,
            "hist_abgeschlossene_bestellungen_anzahl": count,
            "hist_bestellwert_mittel": float(np.mean(amounts)) if count else np.nan,
            "hist_bestellwert_median": float(np.median(amounts)) if count else np.nan,
            "hist_zahlungsabweichung_mittel_tage": float(np.mean(payment_delta)) if count else np.nan,
            "hist_lieferzeit_mittel_tage": float(np.mean(delivery_time)) if count else np.nan,
            "hist_lieferzeit_median_tage": float(np.median(delivery_time)) if count else np.nan,
            "hist_anteil_vorzeitig_oder_puenktlich": ontime / count if count else np.nan,
            "hist_skontonutzung_anteil": discount / count if count else np.nan,
            "hist_wareneingangsabweichung_mittel_tage": float(np.mean(receipt_delta)) if count else np.nan,
            "hist_tage_seit_letzter_lieferantenzahlung": int((cutoff - last_payment).days) if last_payment is not None else np.nan,
            "hist_letztes_abschlussdatum": last_completion,
            "lieferant_ohne_ausreichende_historie": int(count < min_history),
        })
    return pd.DataFrame(rows).set_index("_index")


def supplier_open_items(group: pd.DataFrame) -> pd.DataFrame:
    pay_valid = group["__zahlungsdatum_ist"].isna() | group["__zahlungsdatum_ist"].ge(group["bestelldatum"])
    receipt_valid = group["__wareneingang_ist"].isna() | group["__wareneingang_ist"].ge(group["bestelldatum"])
    base = group.loc[pay_valid & receipt_valid]
    rows: list[dict[str, Any]] = []
    for index, row in group.sort_values(["bestelldatum", "bestellungs_id"]).iterrows():
        cutoff = row["bestelldatum"]
        issued = base["bestelldatum"].lt(cutoff)
        unpaid = base["__zahlungsdatum_ist"].isna() | base["__zahlungsdatum_ist"].ge(cutoff)
        unreceived = base["__wareneingang_ist"].isna() | base["__wareneingang_ist"].ge(cutoff)
        open_orders = base.loc[issued & (unpaid | unreceived)]
        rows.append({
            "_index": index,
            "op_offene_bestellungen_anzahl": len(open_orders),
            "op_offener_bestellwert": float(open_orders["bestellwert"].sum()),
            "op_unbezahlte_bestellungen_anzahl": int((issued & unpaid).sum()),
            "op_unbezahlter_bestellwert": float(base.loc[issued & unpaid, "bestellwert"].sum()),
        })
    return pd.DataFrame(rows).set_index("_index")


def build_supplier_features(
    d: dict[str, pd.DataFrame], nm_daily: pd.DataFrame, cfg: dict[str, Any], holidays: set[pd.Timestamp]
) -> pd.DataFrame:
    orders = d["bestellungen"].copy().rename(columns={
        "zahlungsdatum_ist": "__zahlungsdatum_ist",
        "bezahlter_betrag": "__bezahlter_betrag",
        "tatsaechlicher_wareneingang": "__wareneingang_ist",
    })
    suppliers = d["lieferantenstammdaten"].rename(columns={
        "kategorie": "lieferantenkategorie",
        "zahlungsziel_tage": "lieferant_zahlungsziel_tage",
        "skontofrist_tage": "lieferant_skontofrist_tage",
        "skontosatz": "lieferant_skontosatz",
    })
    out = orders.merge(suppliers, on="lieferanten_id", how="left", validate="many_to_one")
    out["berechnungsstichtag"] = out["bestelldatum"]
    add_calendar_features(out, "bestelldatum", "bestellung", holidays)
    add_calendar_features(out, "faelligkeitsdatum", "faelligkeit", holidays)
    add_calendar_features(out, "erwarteter_wareneingang", "erwarteter_we", holidays)
    out["tage_bestellung_bis_erwarteter_we"] = (out["erwarteter_wareneingang"] - out["bestelldatum"]).dt.days
    out["tage_bestellung_bis_faelligkeit"] = (out["faelligkeitsdatum"] - out["bestelldatum"]).dt.days
    out["tage_bis_erwartete_auszahlung"] = out["tage_bestellung_bis_faelligkeit"]
    out["skontodatum"] = out["bestelldatum"] + pd.to_timedelta(out["skontofrist_tage"], unit="D")
    out["erwarteter_auszahlungsbetrag"] = out["bestellwert"] * (1 - out["skontosatz"].fillna(0) / 100)
    out["dq_lieferant_nicht_verknuepft"] = out["lieferantenname"].isna().astype(int)
    history = pd.concat([
        supplier_history_features(group, cfg["min_history"])
        for _, group in out.groupby("lieferanten_id", sort=False)
    ]).sort_index()
    open_items = pd.concat([
        supplier_open_items(group) for _, group in out.groupby("lieferanten_id", sort=False)
    ]).sort_index()
    rolling = historical_transaction_windows(
        out, "lieferanten_id", "bestelldatum", "bestellwert",
        cfg["rolling_windows"], "hist_bestellung",
    ).sort_index()
    out = out.join(history).join(open_items).join(rolling)
    safe_nm_columns = [
        col for col in nm_daily.columns
        if col == "berechnungsstichtag" or not col.startswith("ziel_")
    ]
    out = out.merge(nm_daily[safe_nm_columns], on="berechnungsstichtag", how="left", validate="many_to_one")
    out["ziel_dq_zahlung_vor_bestelldatum"] = (
        out["__zahlungsdatum_ist"].notna() & out["__zahlungsdatum_ist"].lt(out["bestelldatum"])
    ).astype(int)
    out["ziel_dq_wareneingang_vor_bestelldatum"] = (
        out["__wareneingang_ist"].notna() & out["__wareneingang_ist"].lt(out["bestelldatum"])
    ).astype(int)
    return out.sort_values(["bestelldatum", "bestellungs_id"]).reset_index(drop=True)


# ---------------------------------------------------------------------------
# 5. ZIELGRÖSSEN, ZENSIERUNG UND CHRONOLOGISCHER SPLIT
# ---------------------------------------------------------------------------

def add_split_columns(frame: pd.DataFrame, as_of_col: str, cfg: dict[str, Any]) -> pd.DataFrame:
    out = frame.copy()
    out["as_of_date"] = out[as_of_col]
    out["datensatz_rolle"] = np.select(
        [out["as_of_date"].le(cfg["training_ende"]), out["as_of_date"].between(cfg["validierung_start"], cfg["datenstichtag"])],
        ["Training", "Validierung"], default="Ausserhalb",
    )
    out["split_grenze_training_ende"] = cfg["training_ende"]
    out["split_grenze_validierung_start"] = cfg["validierung_start"]
    out["datenstichtag"] = cfg["datenstichtag"]
    out["beobachtungsende"] = pd.to_datetime(np.where(
        out["datensatz_rolle"].eq("Training"), cfg["training_ende"], cfg["datenstichtag"]
    ))
    return out


def finalize_customer_targets(frame: pd.DataFrame, cfg: dict[str, Any]) -> pd.DataFrame:
    out = add_split_columns(frame, "berechnungsstichtag", cfg)
    bad = out["ziel_dq_zahlung_vor_rechnungsdatum"].ne(0)
    observable = (
        out["__zahlungsdatum_ist"].notna()
        & out["__zahlungsdatum_ist"].ge(out["as_of_date"])
        & out["__zahlungsdatum_ist"].le(out["beobachtungsende"])
        & ~bad
    )
    out["ziel_beobachtbar"] = observable.astype(int)
    out["ziel_zensiert"] = (~observable).astype(int)
    out["ziel_zensierungsgrund"] = np.select(
        [
            bad,
            out["__zahlungsdatum_ist"].isna(),
            out["__zahlungsdatum_ist"].gt(out["beobachtungsende"]),
        ],
        [
            "Zeitlich unplausibles Zahlungsdatum vor Rechnungsdatum",
            "Zahlung bis Beobachtungsende nicht beobachtet",
            "Zahlung erst nach Beobachtungsende beobachtet",
        ],
        default="",
    )
    payment_date = out["__zahlungsdatum_ist"].where(observable)
    payment_amount = out["__bezahlter_betrag"].where(observable)
    days_to_payment = (payment_date - out["as_of_date"]).dt.days
    out["ziel_tatsaechliches_zahlungsdatum"] = payment_date
    out["ziel_tage_bis_zahlung"] = days_to_payment
    out["ziel_zahlungsverzoegerung_tage"] = (payment_date - out["faelligkeitsdatum"]).dt.days
    out["ziel_einzahlungsbetrag"] = payment_amount
    for horizon in cfg["forecast_horizons"]:
        values = pd.Series(pd.NA, index=out.index, dtype="Int64")
        values.loc[observable] = days_to_payment.loc[observable].le(horizon).astype(int)
        out[f"ziel_zahlung_innerhalb_{horizon}_tage"] = values
    out["fuer_spaeteres_supervised_learning_verwendbar"] = observable.astype(int)
    return out.drop(columns=["__zahlungsdatum_ist", "__bezahlter_betrag", "__zahlungsstatus"])


def finalize_supplier_targets(frame: pd.DataFrame, cfg: dict[str, Any]) -> pd.DataFrame:
    out = add_split_columns(frame, "berechnungsstichtag", cfg)
    pay_bad = out["ziel_dq_zahlung_vor_bestelldatum"].ne(0)
    receipt_bad = out["ziel_dq_wareneingang_vor_bestelldatum"].ne(0)
    pay_observable = (
        out["__zahlungsdatum_ist"].notna()
        & out["__zahlungsdatum_ist"].ge(out["as_of_date"])
        & out["__zahlungsdatum_ist"].le(out["beobachtungsende"])
        & ~pay_bad
    )
    receipt_observable = (
        out["__wareneingang_ist"].notna()
        & out["__wareneingang_ist"].ge(out["as_of_date"])
        & out["__wareneingang_ist"].le(out["beobachtungsende"])
        & ~receipt_bad
    )
    out["ziel_beobachtbar"] = pay_observable.astype(int)
    out["ziel_wareneingang_beobachtbar"] = receipt_observable.astype(int)
    out["ziel_zensiert"] = (~pay_observable).astype(int)
    out["ziel_zensierungsgrund"] = np.select(
        [pay_bad, out["__zahlungsdatum_ist"].isna(), out["__zahlungsdatum_ist"].gt(out["beobachtungsende"])],
        ["Zeitlich unplausibles Zahlungsdatum", "Zahlung bis Beobachtungsende nicht beobachtet", "Zahlung erst nach Beobachtungsende beobachtet"],
        default="",
    )
    payment_date = out["__zahlungsdatum_ist"].where(pay_observable)
    payment_amount = out["__bezahlter_betrag"].where(pay_observable)
    days_to_payment = (payment_date - out["as_of_date"]).dt.days
    out["ziel_tatsaechliches_auszahlungsdatum"] = payment_date
    out["ziel_tage_bis_auszahlung"] = days_to_payment
    out["ziel_zahlungsabweichung_tage"] = (payment_date - out["faelligkeitsdatum"]).dt.days
    out["ziel_auszahlungsbetrag"] = payment_amount
    out["ziel_tatsaechliches_wareneingangsdatum"] = out["__wareneingang_ist"].where(receipt_observable)
    out["ziel_lieferzeit_tage"] = (
        out["ziel_tatsaechliches_wareneingangsdatum"] - out["bestelldatum"]
    ).dt.days
    for horizon in cfg["forecast_horizons"]:
        values = pd.Series(pd.NA, index=out.index, dtype="Int64")
        values.loc[pay_observable] = days_to_payment.loc[pay_observable].le(horizon).astype(int)
        out[f"ziel_auszahlung_innerhalb_{horizon}_tage"] = values
    out["fuer_spaeteres_supervised_learning_verwendbar"] = pay_observable.astype(int)
    return out.drop(columns=["__zahlungsdatum_ist", "__bezahlter_betrag", "__wareneingang_ist"])


# ---------------------------------------------------------------------------
# 6. TÄGLICHE UND WÖCHENTLICHE CASHFLOW-ZEITREIHEN
# ---------------------------------------------------------------------------

def add_lags_and_rolling(frame: pd.DataFrame, target_columns: list[str], suffix: str) -> pd.DataFrame:
    out = frame.copy()
    for target in target_columns:
        base = target.removeprefix("ziel_")
        for lag in [1, 7, 14, 30]:
            out[f"{base}_lag_{lag}{suffix}"] = out[target].shift(lag)
        shifted = out[target].shift(1)
        for window in [7, 14, 30]:
            out[f"{base}_summe_letzte_{window}{suffix}"] = shifted.rolling(window, min_periods=1).sum()
            out[f"{base}_mittel_letzte_{window}{suffix}"] = shifted.rolling(window, min_periods=1).mean()
            out[f"{base}_std_letzte_{window}{suffix}"] = shifted.rolling(window, min_periods=1).std(ddof=0)
    return out


def build_daily_cashflow(
    d: dict[str, pd.DataFrame], nm_daily: pd.DataFrame, cfg: dict[str, Any], holidays: set[pd.Timestamp]
) -> pd.DataFrame:
    invoices = d["ausgangsrechnungen"]
    orders = d["bestellungen"]
    history_start = cfg["historie_von"] or min(invoices["rechnungsdatum"].min(), orders["bestelldatum"].min())
    dates = pd.date_range(history_start, cfg["datenstichtag"], freq="D")
    customer_events = invoices.loc[invoices["zahlungsdatum_ist"].between(history_start, cfg["datenstichtag"])]
    supplier_events = orders.loc[orders["zahlungsdatum_ist"].between(history_start, cfg["datenstichtag"])]
    daily = pd.DataFrame({"berechnungsstichtag": dates})
    daily["ziel_kundeneinzahlungen"] = customer_events.groupby("zahlungsdatum_ist")["bezahlter_betrag"].sum().reindex(dates, fill_value=0.0).to_numpy()
    daily["ziel_beschaffungsauszahlungen"] = supplier_events.groupby("zahlungsdatum_ist")["bezahlter_betrag"].sum().reindex(dates, fill_value=0.0).to_numpy()
    daily["ziel_anzahl_kundenzahlungen"] = customer_events.groupby("zahlungsdatum_ist").size().reindex(dates, fill_value=0).to_numpy()
    daily["ziel_anzahl_beschaffungsauszahlungen"] = supplier_events.groupby("zahlungsdatum_ist").size().reindex(dates, fill_value=0).to_numpy()
    daily["ziel_netto_cashflow"] = daily["ziel_kundeneinzahlungen"] - daily["ziel_beschaffungsauszahlungen"]
    add_calendar_features(daily, "berechnungsstichtag", "tag", holidays)
    daily["hist_verfuegbare_tage"] = np.arange(len(daily))
    daily = add_lags_and_rolling(
        daily,
        ["ziel_kundeneinzahlungen", "ziel_beschaffungsauszahlungen", "ziel_netto_cashflow"],
        "t",
    )
    daily = daily.merge(nm_daily, on="berechnungsstichtag", how="left", validate="one_to_one")
    daily = add_split_columns(daily, "berechnungsstichtag", cfg)
    daily["ziel_beobachtbar"] = 1
    daily["ziel_zensiert"] = 0
    daily["ziel_zensierungsgrund"] = ""
    daily["fuer_spaeteres_supervised_learning_verwendbar"] = 1
    return daily


def build_weekly_cashflow(daily: pd.DataFrame, cfg: dict[str, Any]) -> pd.DataFrame:
    base = daily.copy()
    offset = (base["berechnungsstichtag"] - cfg["validierung_start"]).dt.days.mod(7)
    base["periodenbeginn"] = base["berechnungsstichtag"] - pd.to_timedelta(offset, unit="D")
    weekly = base.groupby("periodenbeginn", as_index=False).agg(
        periodenbeobachtung_von=("berechnungsstichtag", "min"),
        periodenbeobachtung_bis=("berechnungsstichtag", "max"),
        tage_enthalten=("berechnungsstichtag", "count"),
        ziel_kundeneinzahlungen=("ziel_kundeneinzahlungen", "sum"),
        ziel_beschaffungsauszahlungen=("ziel_beschaffungsauszahlungen", "sum"),
        ziel_netto_cashflow=("ziel_netto_cashflow", "sum"),
        ziel_anzahl_kundenzahlungen=("ziel_anzahl_kundenzahlungen", "sum"),
        ziel_anzahl_beschaffungsauszahlungen=("ziel_anzahl_beschaffungsauszahlungen", "sum"),
        ziel_nm_verkaufsvolumen_woche=("ziel_nm_verkaufsvolumen_tag", "sum"),
        ziel_nm_bestellvolumen_woche=("ziel_nm_bestellvolumen_tag", "sum"),
    )
    weekly["periodenende"] = weekly["periodenbeginn"] + pd.Timedelta(days=6)
    weekly["as_of_date"] = weekly["periodenbeginn"]
    weekly["periodenjahr"] = weekly["periodenbeginn"].dt.year
    weekly["periodenmonat"] = weekly["periodenbeginn"].dt.month
    weekly["periodenquartal"] = weekly["periodenbeginn"].dt.quarter
    weekly["woche_vollstaendig"] = (
        weekly["tage_enthalten"].eq(7)
        & weekly["periodenbeobachtung_von"].eq(weekly["periodenbeginn"])
        & weekly["periodenbeobachtung_bis"].eq(weekly["periodenende"])
    ).astype(int)
    weekly["hist_verfuegbare_wochen"] = np.arange(len(weekly))
    weekly = add_lags_and_rolling(
        weekly,
        ["ziel_kundeneinzahlungen", "ziel_beschaffungsauszahlungen", "ziel_netto_cashflow"],
        "w",
    )
    weekly["datensatz_rolle"] = np.select(
        [weekly["periodenende"].le(cfg["training_ende"]), weekly["periodenbeginn"].ge(cfg["validierung_start"]) & weekly["periodenende"].le(cfg["datenstichtag"])],
        ["Training", "Validierung"], default="Ausserhalb",
    )
    weekly["split_grenze_training_ende"] = cfg["training_ende"]
    weekly["split_grenze_validierung_start"] = cfg["validierung_start"]
    weekly["datenstichtag"] = cfg["datenstichtag"]
    weekly["beobachtungsende"] = pd.to_datetime(np.where(
        weekly["datensatz_rolle"].eq("Training"), cfg["training_ende"], cfg["datenstichtag"]
    ))
    weekly["ziel_beobachtbar"] = weekly["woche_vollstaendig"]
    weekly["ziel_zensiert"] = weekly["woche_vollstaendig"].eq(0).astype(int)
    weekly["ziel_zensierungsgrund"] = np.where(weekly["woche_vollstaendig"].eq(0), "Unvollstaendige Randwoche", "")
    weekly["fuer_spaeteres_supervised_learning_verwendbar"] = weekly["woche_vollstaendig"]
    return weekly.sort_values("periodenbeginn").reset_index(drop=True)


# ---------------------------------------------------------------------------
# 7. KATALOGE, PRÜFUNGEN UND MODELLSPALTEN
# ---------------------------------------------------------------------------

def build_stichtag_mahnquote(d: dict[str, pd.DataFrame], cfg: dict[str, Any]) -> pd.DataFrame:
    invoices = d["ausgangsrechnungen"].loc[
        d["ausgangsrechnungen"]["rechnungsdatum"].le(cfg["datenstichtag"])
    ].copy()
    invoices["ist_unbezahlt"] = invoices["zahlungsdatum_ist"].isna()
    invoices["ist_mahnung"] = invoices["zahlungsstatus"].astype(str).str.contains("Mahn", case=False, na=False)
    result = invoices.groupby("kunden_id").agg(
        rechnungen_gesamt=("rechnungs_id", "count"),
        unbezahlte_rechnungen=("ist_unbezahlt", "sum"),
        rechnungen_in_mahnstufe=("ist_mahnung", "sum"),
    ).reset_index()
    result["stichtag_mahnquote_alle_rechnungen"] = result["rechnungen_in_mahnstufe"] / result["rechnungen_gesamt"]
    result["stichtag_mahnquote_unbezahlte_rechnungen"] = result["rechnungen_in_mahnstufe"].div(result["unbezahlte_rechnungen"].replace(0, np.nan))
    result["stichtag"] = cfg["datenstichtag"]
    result["als_historisches_modellmerkmal_freigegeben"] = 0
    return result


def model_feature_columns(frame: pd.DataFrame, id_columns: set[str]) -> list[str]:
    excluded_exact = {
        "datensatz_rolle", "as_of_date", "beobachtungsende", "datenstichtag",
        "split_grenze_training_ende", "split_grenze_validierung_start",
        "ziel_zensierungsgrund", "fuer_spaeteres_supervised_learning_verwendbar",
    } | id_columns
    return [
        col for col in frame.columns
        if not col.startswith("ziel_")
        and col not in excluded_exact
        and not col.startswith("dq_")
        and not col.startswith("op_dq_")
        and not col.startswith("__")
    ]


def build_catalog(frames: dict[str, pd.DataFrame]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for dataset, frame in frames.items():
        for column in frame.columns:
            if column.startswith("ziel_"):
                role = "Zielgroesse/Prueffeld"
                formula = "Ergebnisfeld; bei Zensierung leer und niemals Modellinput."
                approved = "Nein"
                leakage = "Durch ziel_-Praefix technisch von Eingangsmerkmalen getrennt."
            elif column.startswith("hist_"):
                role = "Eingangsmerkmal"
                formula = "Ausschliesslich Ereignisse mit Ereignisdatum strikt vor as_of_date."
                approved = "Ja"
                leakage = "Same-Day- und Zukunftsinformation sind ausgeschlossen."
            elif "_lag_" in column or "_letzte_" in column:
                role = "Eingangsmerkmal"
                formula = "Zeitfenster/Lag nach shift(1); aktuelle Periode ist ausgeschlossen."
                approved = "Ja"
                leakage = "Aktueller Zielwert geht nicht in das Merkmal ein."
            elif column.startswith("dq_") or column.startswith("op_dq_"):
                role = "Datenqualitaetsfeld"
                formula = "Regelbasierte Datenqualitaetskennzeichnung."
                approved = "Nein"
                leakage = "Nur fuer Pruefung, nicht als Modellinput vorgesehen."
            elif column in {"datensatz_rolle", "as_of_date", "beobachtungsende", "datenstichtag", "split_grenze_training_ende", "split_grenze_validierung_start"}:
                role = "Split-/Zeitfeld"
                formula = "Aus konfiguriertem Datenstichtag und Validierungsdauer."
                approved = "Nein"
                leakage = "Dient nur der reproduzierbaren Zeittrennung."
            else:
                role = "Eingangsmerkmal/Identifikation"
                formula = "Direktes Quellfeld oder deterministische Kalender-/Betragsableitung."
                approved = "Ja"
                leakage = "Zum jeweiligen Berechnungsstichtag bekannt."
            rows.append({
                "Datensatz": dataset, "Kennzahlenname": column, "Rolle": role,
                "Fachliche_Bedeutung": column.replace("_", " "),
                "Berechnungsregel": formula, "Als_Modellinput_freigegeben": approved,
                "Leakage_Begruendung": leakage,
            })
    return pd.DataFrame(rows)


def run_final_checks(
    d: dict[str, pd.DataFrame], customer: pd.DataFrame, supplier: pd.DataFrame,
    daily: pd.DataFrame, weekly: pd.DataFrame, cfg: dict[str, Any],
    checks: list[dict[str, Any]], model_columns: dict[str, list[str]],
) -> pd.DataFrame:
    frames = {
        "Kundeneinzahlungen": customer,
        "Beschaffungsauszahlungen": supplier,
        "Cashflow_Taeglich": daily,
        "Cashflow_Woechentlich": weekly,
    }
    for dataset, frame in frames.items():
        numeric = frame.select_dtypes(include=[np.number])
        inf_count = int(np.isinf(numeric.to_numpy(dtype=float, na_value=np.nan)).sum()) if len(numeric.columns) else 0
        add_check(checks, "Technik", dataset, "Keine unendlichen Werte", "Divisionen durch 0 müssen leer bleiben.", inf_count)

    key_specs = [
        ("Kundeneinzahlungen", customer, "rechnungs_id"),
        ("Beschaffungsauszahlungen", supplier, "bestellungs_id"),
        ("Cashflow_Taeglich", daily, "as_of_date"),
        ("Cashflow_Woechentlich", weekly, "periodenbeginn"),
    ]
    for dataset, frame, key in key_specs:
        duplicate = frame[key].duplicated(keep=False)
        add_check(checks, "Schluessel", dataset, f"Keine doppelten Ergebnis-Schluessel {key}", "Eine Ergebniszeile je fachlichem Schlüssel.", int(duplicate.sum()), frame.loc[duplicate, key])

    for dataset, frame in [("Kundeneinzahlungen", customer), ("Beschaffungsauszahlungen", supplier), ("Cashflow_Taeglich", daily)]:
        train_dates = set(frame.loc[frame["datensatz_rolle"].eq("Training"), "as_of_date"])
        valid_dates = set(frame.loc[frame["datensatz_rolle"].eq("Validierung"), "as_of_date"])
        overlap = train_dates.intersection(valid_dates)
        add_check(checks, "Split", dataset, "Keine zeitliche Ueberschneidung", "Training und Validierung haben disjunkte Stichtage.", len(overlap), overlap)

    validation_daily = daily.loc[daily["datensatz_rolle"].eq("Validierung")]
    validation_weekly = weekly.loc[weekly["datensatz_rolle"].eq("Validierung")]
    add_check(checks, "Split", "Cashflow_Taeglich", "Exakt konfigurierte Validierungstage", f"Validierung umfasst {cfg['validierungstage']} Kalendertage.", abs(validation_daily["as_of_date"].nunique() - cfg["validierungstage"]))
    expected_weeks = cfg["validierungstage"] // 7
    add_check(checks, "Split", "Cashflow_Woechentlich", "Vollstaendige Validierungswochen", f"Bei durch 7 teilbarer Dauer werden {expected_weeks} vollständige Wochen erwartet.", abs(len(validation_weekly) - expected_weeks) if cfg["validierungstage"] % 7 == 0 else 0)
    incomplete_validation = int(validation_weekly["woche_vollstaendig"].ne(1).sum())
    add_check(checks, "Split", "Cashflow_Woechentlich", "Keine unvollstaendige Validierungswoche", "Jede Validierungswoche enthält sieben Tage.", incomplete_validation, validation_weekly.loc[validation_weekly["woche_vollstaendig"].ne(1), "periodenbeginn"])

    row_checks = [
        ("Kundeneinzahlungen", len(d["ausgangsrechnungen"]), len(customer)),
        ("Beschaffungsauszahlungen", len(d["bestellungen"]), len(supplier)),
    ]
    for dataset, expected, actual in row_checks:
        add_check(checks, "Bestandsabgleich", dataset, "Keine Zeilenvervielfachung oder Loeschung", f"Quelle={expected}, Ergebnis={actual}.", abs(expected - actual))

    customer_hist_bad = customer["hist_letztes_zahlungsdatum"].notna() & customer["hist_letztes_zahlungsdatum"].ge(customer["as_of_date"])
    supplier_hist_bad = supplier["hist_letztes_abschlussdatum"].notna() & supplier["hist_letztes_abschlussdatum"].ge(supplier["as_of_date"])
    add_check(checks, "Data Leakage", "Kundeneinzahlungen", "Historie strikt vor Stichtag", "hist_letztes_zahlungsdatum < as_of_date.", int(customer_hist_bad.sum()), customer.loc[customer_hist_bad, "rechnungs_id"])
    add_check(checks, "Data Leakage", "Beschaffungsauszahlungen", "Historie strikt vor Stichtag", "hist_letztes_abschlussdatum < as_of_date.", int(supplier_hist_bad.sum()), supplier.loc[supplier_hist_bad, "bestellungs_id"])

    daily_lag_expected = daily["ziel_netto_cashflow"].shift(1)
    daily_lag_bad = (daily["netto_cashflow_lag_1t"] - daily_lag_expected).abs().gt(0.005)
    weekly_lag_expected = weekly["ziel_netto_cashflow"].shift(1)
    weekly_lag_bad = (weekly["netto_cashflow_lag_1w"] - weekly_lag_expected).abs().gt(0.005)
    add_check(checks, "Data Leakage", "Cashflow_Taeglich", "Lag 1 entspricht Vortag", "Lag1[t] = Ziel[t-1].", int(daily_lag_bad.sum()))
    add_check(checks, "Data Leakage", "Cashflow_Woechentlich", "Lag 1 entspricht Vorwoche", "Lag1[w] = Ziel[w-1].", int(weekly_lag_bad.sum()))

    target_in_features = [
        f"{dataset}:{column}" for dataset, columns in model_columns.items()
        for column in columns if column.startswith("ziel_")
    ]
    add_check(checks, "Data Leakage", "Alle", "Keine Zielspalten in Modellmerkmalen", "Kein freigegebenes Modellmerkmal beginnt mit ziel_.", len(target_in_features), target_in_features)

    customer_unmasked = customer["ziel_zensiert"].eq(1) & customer[[
        "ziel_tatsaechliches_zahlungsdatum", "ziel_tage_bis_zahlung", "ziel_einzahlungsbetrag"
    ]].notna().any(axis=1)
    supplier_unmasked = supplier["ziel_zensiert"].eq(1) & supplier[[
        "ziel_tatsaechliches_auszahlungsdatum", "ziel_tage_bis_auszahlung", "ziel_auszahlungsbetrag"
    ]].notna().any(axis=1)
    add_check(checks, "Data Leakage", "Kundeneinzahlungen", "Zensierte Ziele sind leer", "Zukunftszielwerte werden nicht nachträglich ergänzt.", int(customer_unmasked.sum()), customer.loc[customer_unmasked, "rechnungs_id"])
    add_check(checks, "Data Leakage", "Beschaffungsauszahlungen", "Zensierte Ziele sind leer", "Zukunftszielwerte werden nicht nachträglich ergänzt.", int(supplier_unmasked.sum()), supplier.loc[supplier_unmasked, "bestellungs_id"])

    history_start = cfg["historie_von"] or daily["berechnungsstichtag"].min()
    raw_in = float(d["ausgangsrechnungen"].loc[d["ausgangsrechnungen"]["zahlungsdatum_ist"].between(history_start, cfg["datenstichtag"]), "bezahlter_betrag"].sum())
    raw_out = float(d["bestellungen"].loc[d["bestellungen"]["zahlungsdatum_ist"].between(history_start, cfg["datenstichtag"]), "bezahlter_betrag"].sum())
    for name, expected, actual in [
        ("Kundeneinzahlungen", raw_in, daily["ziel_kundeneinzahlungen"].sum()),
        ("Beschaffungsauszahlungen", raw_out, daily["ziel_beschaffungsauszahlungen"].sum()),
        ("Netto-Cashflow", daily["ziel_kundeneinzahlungen"].sum() - daily["ziel_beschaffungsauszahlungen"].sum(), daily["ziel_netto_cashflow"].sum()),
    ]:
        delta = abs(float(expected) - float(actual))
        add_check(
            checks, "Kontrollsumme", name, f"Kontrollsumme {name}",
            f"Rohdatensumme muss der Tagesreihe entsprechen; Differenz={delta:.6f}.",
            int(delta >= 0.005),
        )
    return pd.DataFrame(checks)


def dummy_controls(
    d: dict[str, pd.DataFrame], customer: pd.DataFrame, supplier: pd.DataFrame,
    cfg: dict[str, Any]
) -> pd.DataFrame:
    invoices = d["ausgangsrechnungen"]
    plausible = invoices.loc[
        invoices["zahlungsdatum_ist"].notna()
        & invoices["zahlungsdatum_ist"].ge(invoices["rechnungsdatum"])
    ].copy()
    plausible["delay"] = (plausible["zahlungsdatum_ist"] - plausible["faelligkeitsdatum"]).dt.days
    actual = {
        "Kundenrechnungen": len(invoices),
        "Beschaffungsvorgaenge": len(d["bestellungen"]),
        "Kunden_Training": int(customer["datensatz_rolle"].eq("Training").sum()),
        "Kunden_Validierung": int(customer["datensatz_rolle"].eq("Validierung").sum()),
        "Beschaffung_Training": int(supplier["datensatz_rolle"].eq("Training").sum()),
        "Beschaffung_Validierung": int(supplier["datensatz_rolle"].eq("Validierung").sum()),
        "Plausible_Zahlungen": len(plausible),
        "Unplausible_Zahlungen": int((invoices["zahlungsdatum_ist"].notna() & invoices["zahlungsdatum_ist"].lt(invoices["rechnungsdatum"])).sum()),
        "Mittlerer_Zahlungsverzug": float(plausible["delay"].mean()),
        "Verspaetungsquote": float(plausible["delay"].gt(0).mean()),
    }
    expected = {
        "Kundenrechnungen": 20902, "Beschaffungsvorgaenge": 1066,
        "Kunden_Training": 20078, "Kunden_Validierung": 824,
        "Beschaffung_Training": 1003, "Beschaffung_Validierung": 63,
        "Plausible_Zahlungen": 19848, "Unplausible_Zahlungen": 374,
        "Mittlerer_Zahlungsverzug": 1.45, "Verspaetungsquote": 0.548,
    }
    rows = []
    for key, expected_value in expected.items():
        actual_value = actual[key]
        tolerance = 0.02 if key in {"Mittlerer_Zahlungsverzug", "Verspaetungsquote"} else 0
        ok = abs(float(actual_value) - float(expected_value)) <= tolerance
        rows.append({
            "Kontrollwert": key, "Soll_nur_Dummy": expected_value, "Ist": actual_value,
            "Status": "PASS" if ok else "WARN",
            "Hinweis": "Nur Plausibilitaetswarnung im Dummy-Modus; keine Bedingung fuer Echtdaten.",
        })
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# 8. EXPORT: CSV, EXCEL, GRAFIKEN, BERICHT UND PROTOKOLL
# ---------------------------------------------------------------------------

def write_csv(frame: pd.DataFrame, path: Path) -> None:
    frame.to_csv(
        path, index=False, sep=";", decimal=",", encoding="utf-8-sig",
        date_format="%Y-%m-%d", na_rep="",
    )


def split_frame(frame: pd.DataFrame, role: str, sort_columns: list[str]) -> pd.DataFrame:
    return frame.loc[frame["datensatz_rolle"].eq(role)].sort_values(sort_columns).reset_index(drop=True)


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        if sheet.max_row >= 1 and sheet.max_column >= 1:
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column_cells in sheet.iter_cols(min_row=1, max_row=min(sheet.max_row, 250)):
            max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 10), 34)
    workbook.save(path)


def write_excel_files(
    output_folder: Path, outputs: dict[str, pd.DataFrame], catalog: pd.DataFrame,
    provenance: pd.DataFrame, split_summary: pd.DataFrame, checks: pd.DataFrame,
    source_inventory: pd.DataFrame, context_inventory: pd.DataFrame,
    dummy_check: pd.DataFrame, mahnquote: pd.DataFrame, cfg: dict[str, Any],
) -> tuple[Path, Path]:
    main_path = output_folder / "11_CASH_AI_Aufbereitete_Datensaetze.xlsx"
    readme = pd.DataFrame([
        ["Zweck", "Leakage-sichere Datensaetze fuer ein spaeteres Cashflow-Prognosemodell; noch kein Modelltraining."],
        ["Datenmodus", cfg["data_mode"]],
        ["Datenstichtag", cfg["datenstichtag"]],
        ["Training bis", cfg["training_ende"]],
        ["Validierung von", cfg["validierung_start"]],
        ["Validierung bis", cfg["datenstichtag"]],
        ["Wichtig", "Zensierte Zielwerte sind leer. Nur fuer_spaeteres_supervised_learning_verwendbar=1 trainieren."],
        ["CSV", "Vollstaendige Modelltabellen; UTF-8-BOM, Semikolon, ISO-Datum."],
    ], columns=["Thema", "Erlaeuterung"])
    with pd.ExcelWriter(main_path, engine="openpyxl", date_format="YYYY-MM-DD", datetime_format="YYYY-MM-DD") as writer:
        readme.to_excel(writer, sheet_name="README", index=False)
        outputs["customer_total"].to_excel(writer, sheet_name="Kunden_Gesamt", index=False)
        outputs["customer_train"].to_excel(writer, sheet_name="Kunden_Training", index=False)
        outputs["customer_valid"].to_excel(writer, sheet_name="Kunden_Validierung", index=False)
        outputs["supplier_total"].to_excel(writer, sheet_name="Beschaffung_Gesamt", index=False)
        outputs["supplier_train"].to_excel(writer, sheet_name="Beschaffung_Training", index=False)
        outputs["supplier_valid"].to_excel(writer, sheet_name="Beschaffung_Validierung", index=False)
        outputs["daily_total"].to_excel(writer, sheet_name="Cashflow_Taeglich", index=False)
        outputs["weekly_total"].to_excel(writer, sheet_name="Cashflow_Woechentlich", index=False)
        catalog.to_excel(writer, sheet_name="Kennzahlenkatalog", index=False)
        provenance.to_excel(writer, sheet_name="Spaltenherkunft", index=False)
        split_summary.to_excel(writer, sheet_name="Split_Dokumentation", index=False)
        checks.to_excel(writer, sheet_name="Datenqualitaet", index=False)
        checks.loc[checks["Kategorie"].eq("Data Leakage")].to_excel(writer, sheet_name="Leakage_Pruefung", index=False)
        mahnquote.to_excel(writer, sheet_name="Stichtag_Mahnquote", index=False)
    style_workbook(main_path)

    quality_path = output_folder / "12_CASH_AI_Datenqualitaetsbericht.xlsx"
    with pd.ExcelWriter(quality_path, engine="openpyxl", date_format="YYYY-MM-DD") as writer:
        checks.to_excel(writer, sheet_name="Pruefungen", index=False)
        checks.loc[checks["Status"].isin(["WARN", "FAIL"])].to_excel(writer, sheet_name="Auffaelligkeiten", index=False)
        source_inventory.to_excel(writer, sheet_name="Quellinventar", index=False)
        context_inventory.to_excel(writer, sheet_name="Dateiinventar", index=False)
        dummy_check.to_excel(writer, sheet_name="Dummy_Kontrollwerte", index=False)
    style_workbook(quality_path)
    return main_path, quality_path


def create_charts(
    output_folder: Path, daily: pd.DataFrame, customer: pd.DataFrame,
    supplier: pd.DataFrame, nm_daily: pd.DataFrame,
) -> list[Path]:
    chart_folder = output_folder / "Grafiken"
    chart_folder.mkdir(exist_ok=True)
    sns.set_theme(style="whitegrid")
    paths: list[Path] = []

    def save(name: str) -> None:
        path = chart_folder / name
        plt.tight_layout()
        plt.savefig(path, dpi=180, bbox_inches="tight")
        plt.close()
        paths.append(path)

    monthly = daily.set_index("berechnungsstichtag")[[
        "ziel_kundeneinzahlungen", "ziel_beschaffungsauszahlungen", "ziel_netto_cashflow"
    ]].resample("MS").sum()
    for column, title, filename, color in [
        ("ziel_kundeneinzahlungen", "Kundeneinzahlungen im Zeitverlauf", "01_Kundeneinzahlungen.png", "#2E75B6"),
        ("ziel_beschaffungsauszahlungen", "Beschaffungsauszahlungen im Zeitverlauf", "02_Beschaffungsauszahlungen.png", "#C55A11"),
        ("ziel_netto_cashflow", "Netto-Cashflow im Zeitverlauf", "03_Netto_Cashflow.png", "#548235"),
    ]:
        plt.figure(figsize=(11, 4.8))
        plt.plot(monthly.index, monthly[column], color=color, linewidth=2)
        plt.axhline(0, color="black", linewidth=0.8)
        plt.title(title)
        plt.ylabel("EUR pro Monat")
        save(filename)

    observed_customer = customer.loc[customer["ziel_beobachtbar"].eq(1)].copy()
    observed_customer["monat"] = observed_customer["as_of_date"].dt.to_period("M").dt.to_timestamp()
    delay_month = observed_customer.groupby("monat")["ziel_zahlungsverzoegerung_tage"].mean()
    plt.figure(figsize=(11, 4.8))
    plt.plot(delay_month.index, delay_month, color="#7030A0", linewidth=2)
    plt.axhline(0, color="black", linewidth=0.8)
    plt.title("Durchschnittliche Abweichung vom Zahlungsziel")
    plt.ylabel("Tage")
    save("04_Zahlungszielabweichung.png")

    segment = observed_customer.assign(_late=observed_customer["ziel_zahlungsverzoegerung_tage"].gt(0)).groupby("segment")["_late"].mean().sort_values()
    plt.figure(figsize=(8, 4.8))
    sns.barplot(x=segment.index, y=segment.values, color="#5B9BD5")
    plt.title("Verspaetungsquote nach Kundensegment")
    plt.ylabel("Anteil")
    plt.ylim(0, 1)
    save("05_Verspaetungsquote_Segment.png")

    nm_month = nm_daily.set_index("berechnungsstichtag")[[
        "ziel_nm_verkaufsvolumen_tag", "ziel_nm_bestellvolumen_tag"
    ]].resample("MS").sum()
    plt.figure(figsize=(11, 4.8))
    plt.plot(nm_month.index, nm_month.iloc[:, 0], label="Verkaufsvolumen", linewidth=2)
    plt.plot(nm_month.index, nm_month.iloc[:, 1], label="Bestellvolumen", linewidth=2)
    plt.title("Neumaschinen: Verkaufs- und Bestellvolumen")
    plt.ylabel("EUR pro Monat")
    plt.legend()
    save("06_Neumaschinen_Verkauf_Bestellung.png")

    split_counts = pd.DataFrame({
        "Kundeneinzahlungen": customer["datensatz_rolle"].value_counts(),
        "Beschaffungsauszahlungen": supplier["datensatz_rolle"].value_counts(),
    }).fillna(0).T
    split_counts.plot(kind="bar", figsize=(9, 5), color=["#4472C4", "#ED7D31", "#A5A5A5"])
    plt.title("Verteilung Training und Validierung")
    plt.ylabel("Zeilen")
    plt.xticks(rotation=0)
    save("07_Trainings_Validierungsverteilung.png")
    return paths


def build_split_summary(outputs: dict[str, pd.DataFrame]) -> pd.DataFrame:
    rows = []
    specs = [
        ("Kundeneinzahlungen", outputs["customer_train"], "Training", "rechnungsbetrag"),
        ("Kundeneinzahlungen", outputs["customer_valid"], "Validierung", "rechnungsbetrag"),
        ("Beschaffungsauszahlungen", outputs["supplier_train"], "Training", "bestellwert"),
        ("Beschaffungsauszahlungen", outputs["supplier_valid"], "Validierung", "bestellwert"),
        ("Cashflow taeglich", outputs["daily_train"], "Training", "ziel_netto_cashflow"),
        ("Cashflow taeglich", outputs["daily_valid"], "Validierung", "ziel_netto_cashflow"),
        ("Cashflow woechentlich", outputs["weekly_train"], "Training", "ziel_netto_cashflow"),
        ("Cashflow woechentlich", outputs["weekly_valid"], "Validierung", "ziel_netto_cashflow"),
    ]
    for dataset, frame, role, amount in specs:
        date_col = "periodenbeginn" if "woechentlich" in dataset else "as_of_date"
        rows.append({
            "Datensatz": dataset, "Rolle": role, "Zeilen": len(frame),
            "Von": frame[date_col].min() if len(frame) else pd.NaT,
            "Bis": frame[date_col].max() if len(frame) else pd.NaT,
            "Beobachtbare_Ziele": int(frame["ziel_beobachtbar"].fillna(0).sum()),
            "Zensierte_Ziele": int(frame["ziel_zensiert"].fillna(0).sum()),
            "Relevante_Summe": float(frame[amount].sum()),
        })
    return pd.DataFrame(rows)


def create_markdown_report(
    output_folder: Path, cfg: dict[str, Any], source_inventory: pd.DataFrame,
    outputs: dict[str, pd.DataFrame], checks: pd.DataFrame, dummy_check: pd.DataFrame,
    charts: list[Path], model_columns: dict[str, list[str]],
) -> Path:
    path = output_folder / "13_CASH_AI_Kurzbericht.md"
    customer = outputs["customer_total"]
    supplier = outputs["supplier_total"]
    failures = int(checks["Status"].eq("FAIL").sum())
    warnings = int(checks["Status"].eq("WARN").sum())
    lines = [
        "# CASH-AI – Kurzbericht zur Datenaufbereitung",
        "",
        f"**Datenmodus:** {cfg['data_mode']}  ",
        f"**Datenstichtag:** {cfg['datenstichtag']:%Y-%m-%d}  ",
        f"**Training:** bis {cfg['training_ende']:%Y-%m-%d}  ",
        f"**Validierung:** {cfg['validierung_start']:%Y-%m-%d} bis {cfg['datenstichtag']:%Y-%m-%d} ({cfg['validierungstage']} Tage)  ",
        "**Modelltraining:** nicht Bestandteil dieses Programmlaufs.",
        "",
        "## Verwendete Quellen",
        "",
        "| Datensatz | Datei | Zeilen | Spalten |",
        "|---|---|---:|---:|",
    ]
    for _, row in source_inventory.iterrows():
        lines.append(f"| {row['Datensatz']} | {row['Quelldatei']} | {row['Zeilen']} | {row['Spalten']} |")
    lines += [
        "",
        "## Ergebnisumfang",
        "",
        f"- Kundeneinzahlungen: {len(customer):,} Rechnungszeilen, davon {int(customer['datensatz_rolle'].eq('Training').sum()):,} Training und {int(customer['datensatz_rolle'].eq('Validierung').sum()):,} Validierung.",
        f"- Beschaffungsauszahlungen: {len(supplier):,} Bestellzeilen, davon {int(supplier['datensatz_rolle'].eq('Training').sum()):,} Training und {int(supplier['datensatz_rolle'].eq('Validierung').sum()):,} Validierung.",
        f"- Freigegebene Kundenmerkmale: {len(model_columns['Kundeneinzahlungen'])}; Beschaffungsmerkmale: {len(model_columns['Beschaffungsauszahlungen'])}.",
        f"- Beobachtbare/zensierte Kundenziele: {int(customer['ziel_beobachtbar'].sum()):,}/{int(customer['ziel_zensiert'].sum()):,}.",
        f"- Beobachtbare/zensierte Beschaffungsziele: {int(supplier['ziel_beobachtbar'].sum()):,}/{int(supplier['ziel_zensiert'].sum()):,}.",
        "",
        "## Data-Leakage-Regel",
        "",
        "Historische Kennzahlen verwenden nur Ereignisse mit Ereignisdatum **strikt vor** dem jeweiligen `as_of_date`. Same-Day-Ereignisse werden wegen fehlender Uhrzeit nicht als Historie verwendet. Lag- und Rolling-Merkmale werden vor der Aggregation um eine Periode verschoben. Zielspalten beginnen mit `ziel_`; nicht beobachtbare Zielwerte werden leer ausgegeben und mit `ziel_zensiert = 1` gekennzeichnet.",
        "",
        "## Datenqualität",
        "",
        f"Insgesamt wurden {len(checks)} Kontrollen ausgeführt: {failures} FAIL und {warnings} WARN. WARN-Fälle bleiben transparent erhalten; kritische FAIL-Ergebnisse beenden das Programm.",
        "",
        "## Grafiken",
        "",
    ]
    for chart in charts:
        lines.append(f"![{chart.stem}](Grafiken/{chart.name})")
        lines.append("")
    if cfg["data_mode"] == "dummy":
        lines += [
            "## Dummy-Kontrollwerte",
            "",
            "Die im Prüfbericht enthaltenen Dummy-Werte dienen nur der Reproduzierbarkeitskontrolle und sind keine Bedingungen für Echtdaten.",
            "",
        ]
        for _, row in dummy_check.iterrows():
            lines.append(f"- {row['Kontrollwert']}: Ist {row['Ist']} – {row['Status']}")
    lines += [
        "",
        "## Austausch gegen Echtdaten",
        "",
        "1. Rohdateien in einen neuen Eingabeordner kopieren.",
        "2. `data_mode` auf `real`, `input_folder` und `datenstichtag` anpassen.",
        "3. Bei deutsch formatierten Datumswerten `date_order = dayfirst` setzen.",
        "4. Abweichende Feldnamen ausschließlich in `COLUMN_ALIASES` ergänzen.",
        "5. Programm erneut ausführen und zuerst `12_CASH_AI_Datenqualitaetsbericht.xlsx` prüfen.",
        "6. Für ein späteres Modell nur Zeilen mit `fuer_spaeteres_supervised_learning_verwendbar = 1` und die Spalten aus `15_CASH_AI_Modellspalten.json` verwenden.",
        "",
        "Die Feature-Engineering-Funktionen sind in Dummy- und Realmodus identisch; es gibt keine fest eingebauten Dummy-Zeilenzahlen in der Berechnungslogik.",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def export_results(
    d: dict[str, pd.DataFrame], customer: pd.DataFrame, supplier: pd.DataFrame,
    daily: pd.DataFrame, weekly: pd.DataFrame, nm_daily: pd.DataFrame,
    cfg: dict[str, Any], checks: pd.DataFrame, catalog: pd.DataFrame,
    provenance: pd.DataFrame, source_inventory: pd.DataFrame,
    context_inventory: pd.DataFrame, dummy_check: pd.DataFrame,
    mahnquote: pd.DataFrame, model_columns: dict[str, list[str]],
) -> dict[str, Path]:
    output_folder = cfg["output_folder"]
    outputs = {
        "customer_total": customer.loc[customer["datensatz_rolle"].ne("Ausserhalb")].reset_index(drop=True),
        "customer_train": split_frame(customer, "Training", ["as_of_date", "rechnungs_id"]),
        "customer_valid": split_frame(customer, "Validierung", ["as_of_date", "rechnungs_id"]),
        "supplier_total": supplier.loc[supplier["datensatz_rolle"].ne("Ausserhalb")].reset_index(drop=True),
        "supplier_train": split_frame(supplier, "Training", ["as_of_date", "bestellungs_id"]),
        "supplier_valid": split_frame(supplier, "Validierung", ["as_of_date", "bestellungs_id"]),
        "daily_total": daily.loc[daily["datensatz_rolle"].ne("Ausserhalb")].reset_index(drop=True),
        "daily_train": split_frame(daily, "Training", ["as_of_date"]),
        "daily_valid": split_frame(daily, "Validierung", ["as_of_date"]),
        "weekly_total": weekly.loc[weekly["datensatz_rolle"].ne("Ausserhalb")].reset_index(drop=True),
        "weekly_train": split_frame(weekly, "Training", ["periodenbeginn"]),
        "weekly_valid": split_frame(weekly, "Validierung", ["periodenbeginn"]),
    }
    file_map = {
        "01_Kundeneinzahlungen_Gesamt.csv": outputs["customer_total"],
        "02_Kundeneinzahlungen_Training.csv": outputs["customer_train"],
        "03_Kundeneinzahlungen_Validierung.csv": outputs["customer_valid"],
        "04_Beschaffungsauszahlungen_Gesamt.csv": outputs["supplier_total"],
        "05_Beschaffungsauszahlungen_Training.csv": outputs["supplier_train"],
        "06_Beschaffungsauszahlungen_Validierung.csv": outputs["supplier_valid"],
        "07_Cashflow_Taeglich_Training.csv": outputs["daily_train"],
        "08_Cashflow_Taeglich_Validierung.csv": outputs["daily_valid"],
        "09_Cashflow_Woechentlich_Training.csv": outputs["weekly_train"],
        "10_Cashflow_Woechentlich_Validierung.csv": outputs["weekly_valid"],
    }
    paths: dict[str, Path] = {}
    for filename, frame in file_map.items():
        path = output_folder / filename
        write_csv(frame, path)
        paths[filename] = path
        logging.info("CSV geschrieben: %s (%s Zeilen)", filename, len(frame))

    split_summary = build_split_summary(outputs)
    main_xlsx, quality_xlsx = write_excel_files(
        output_folder, outputs, catalog, provenance, split_summary, checks,
        source_inventory, context_inventory, dummy_check, mahnquote, cfg,
    )
    paths[main_xlsx.name] = main_xlsx
    paths[quality_xlsx.name] = quality_xlsx
    charts = create_charts(output_folder, outputs["daily_total"], outputs["customer_total"], outputs["supplier_total"], nm_daily)
    report_path = create_markdown_report(
        output_folder, cfg, source_inventory, outputs, checks, dummy_check,
        charts, model_columns,
    )
    paths[report_path.name] = report_path
    model_path = output_folder / "15_CASH_AI_Modellspalten.json"
    model_payload = {
        "hinweis": "Noch kein Modelltraining. Nur freigegebene Zeilen und diese Eingangsmerkmale verwenden.",
        "datenstichtag": cfg["datenstichtag"],
        "training_ende": cfg["training_ende"],
        "validierung_start": cfg["validierung_start"],
        "eingangsmerkmale": model_columns,
        "zielspalten_kunden": [col for col in customer.columns if col.startswith("ziel_")],
        "zielspalten_beschaffung": [col for col in supplier.columns if col.startswith("ziel_")],
    }
    model_path.write_text(json.dumps(json_safe(model_payload), ensure_ascii=False, indent=2), encoding="utf-8")
    paths[model_path.name] = model_path
    manifest_path = output_folder / "16_CASH_AI_Manifest.json"
    manifest = {
        "status": "PASS",
        "konfiguration": json_safe(cfg),
        "dateien": {
            name: {"pfad": str(path), "bytes": path.stat().st_size, "sha256": sha256(path)}
            for name, path in paths.items()
        },
        "pruefungen": {
            "gesamt": len(checks), "pass": int(checks["Status"].eq("PASS").sum()),
            "warn": int(checks["Status"].eq("WARN").sum()), "fail": int(checks["Status"].eq("FAIL").sum()),
        },
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    paths[manifest_path.name] = manifest_path
    return paths


# ---------------------------------------------------------------------------
# 9. HAUPTPROGRAMM
# ---------------------------------------------------------------------------

def main() -> None:
    args = parse_args()
    cfg = effective_config(args)
    log_path = setup_logging(cfg["output_folder"])
    logging.info("CASH-AI Datenaufbereitung gestartet; Modus=%s", cfg["data_mode"])
    logging.info("Eingabe=%s | Ausgabe=%s", cfg["input_folder"], cfg["output_folder"])
    logging.info(
        "Training bis %s | Validierung %s bis %s",
        cfg["training_ende"].date(), cfg["validierung_start"].date(), cfg["datenstichtag"].date(),
    )
    checks: list[dict[str, Any]] = []
    try:
        if not cfg["input_folder"].exists():
            raise FileNotFoundError(
                f"Eingabeordner fehlt: {cfg['input_folder']}. "
                "Lege dort die sechs Roh-CSV-Dateien ab."
            )
        context_inventory = inspect_context_files(cfg["input_folder"])
        d, source_inventory, provenance = load_sources(cfg, checks)
        validate_sources(d, checks)
        history_start = cfg["historie_von"] or min(
            d["ausgangsrechnungen"]["rechnungsdatum"].min(), d["bestellungen"]["bestelldatum"].min()
        )
        holidays = set(austrian_holidays(range(history_start.year - 1, cfg["datenstichtag"].year + 2)))
        nm_daily = build_neumaschinen_daily(d, cfg)
        customer = finalize_customer_targets(
            build_customer_features(d, nm_daily, cfg, holidays), cfg
        )
        supplier = finalize_supplier_targets(
            build_supplier_features(d, nm_daily, cfg, holidays), cfg
        )
        daily = build_daily_cashflow(d, nm_daily, cfg, holidays)
        weekly = build_weekly_cashflow(daily, cfg)
        mahnquote = build_stichtag_mahnquote(d, cfg)
        model_columns = {
            "Kundeneinzahlungen": model_feature_columns(customer, {"rechnungs_id", "kunden_id", "rechnungsdatum", "faelligkeitsdatum", "skontodatum", "hist_letztes_zahlungsdatum"}),
            "Beschaffungsauszahlungen": model_feature_columns(supplier, {"bestellungs_id", "lieferanten_id", "bestelldatum", "faelligkeitsdatum", "erwarteter_wareneingang", "skontodatum", "hist_letztes_abschlussdatum"}),
            "Cashflow_Taeglich": model_feature_columns(daily, {"berechnungsstichtag"}),
            "Cashflow_Woechentlich": model_feature_columns(weekly, {"periodenbeginn", "periodenende", "periodenbeobachtung_von", "periodenbeobachtung_bis"}),
        }
        checks_df = run_final_checks(
            d, customer, supplier, daily, weekly, cfg, checks, model_columns
        )
        if checks_df["Status"].eq("FAIL").any():
            failed = checks_df.loc[checks_df["Status"].eq("FAIL"), ["Datensatz", "Pruefung", "Verletzungen"]]
            raise RuntimeError("Kritische Prüfungen fehlgeschlagen:\n" + failed.to_string(index=False))
        catalog = build_catalog({
            "Kundeneinzahlungen": customer,
            "Beschaffungsauszahlungen": supplier,
            "Cashflow_Taeglich": daily,
            "Cashflow_Woechentlich": weekly,
        })
        dummy_check = dummy_controls(d, customer, supplier, cfg) if cfg["data_mode"] == "dummy" else pd.DataFrame(columns=["Kontrollwert", "Soll_nur_Dummy", "Ist", "Status", "Hinweis"])
        paths = export_results(
            d, customer, supplier, daily, weekly, nm_daily, cfg, checks_df,
            catalog, provenance, source_inventory, context_inventory, dummy_check,
            mahnquote, model_columns,
        )
        logging.info("CASH-AI Datenaufbereitung erfolgreich abgeschlossen.")
        print("\nFERTIG – kein Prognosemodell wurde trainiert.")
        print(f"Pruefungen: {int(checks_df['Status'].eq('PASS').sum())} PASS, {int(checks_df['Status'].eq('WARN').sum())} WARN, 0 FAIL")
        print(f"Training: bis {cfg['training_ende']:%Y-%m-%d}")
        print(f"Validierung: {cfg['validierung_start']:%Y-%m-%d} bis {cfg['datenstichtag']:%Y-%m-%d}")
        print("Erzeugte Hauptdateien:")
        for name in sorted(paths):
            print(f"  - {paths[name]}")
        print(f"  - {log_path}")
    except Exception:
        logging.exception("Programmabbruch")
        print(
            f"\nFEHLER: Verarbeitung abgebrochen. Details stehen in:\n{log_path}",
            file=sys.stderr,
        )
        raise


if __name__ == "__main__":
    main()
